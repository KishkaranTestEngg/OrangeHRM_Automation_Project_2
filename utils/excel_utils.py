from openpyxl import load_workbook


def read_login_test_data():
    file_path = "data/OrangeHRM_TestData.xlsx"

    workbook = load_workbook(file_path)
    sheet = workbook["TC001_LoginData"]

    test_data = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        test_data.append({
            "sl_no": row[0],
            "test_id": row[1],
            "tester": row[2],
            "date": row[3],
            "test_parameter": row[4],
            "username": row[5],
            "password": row[6],
            "test_result": row[7]
        })

    workbook.close()

    return test_data